"""Copy each card's Value / Budget / Status from the design spreadsheet into
tests/expected_balance.json, so the Godot tests check the code against the sheet.

Run after changing numbers in docs/design/Dungeon_Card_Set_v1.xlsx
(save the sheet in Excel, LibreOffice or Google Sheets first so formulas
have calculated values):

    pip install openpyxl
    python tools/export_expected_balance.py
"""
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SHEET = ROOT / "docs" / "design" / "Dungeon_Card_Set_v1.xlsx"
OUT = ROOT / "tests" / "expected_balance.json"

ws = load_workbook(SHEET, data_only=True)["Cards"]
header = [c.value for c in ws[1]]
col = {name: header.index(name) for name in ("ID", "Value", "Budget", "Status")}

result = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    card_id = row[col["ID"]]
    if not card_id:
        continue
    if row[col["Value"]] is None:
        raise SystemExit(f"{card_id} has no calculated Value; open and save the sheet first.")
    result[card_id] = {
        "value": row[col["Value"]],
        "budget": row[col["Budget"]],
        "status": row[col["Status"]],
    }

OUT.write_text(json.dumps(result, indent=1) + "\n")
print(f"Wrote {len(result)} cards to {OUT.relative_to(ROOT)}")
